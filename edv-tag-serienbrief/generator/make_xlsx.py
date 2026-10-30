"""Builds Bewerbungsliste.xlsx: the mail-merge data source for the Serienbrief lesson.

Design rules (they are what makes the Word mail merge reliable):
- exactly ONE sheet, headers in row 1, data from row 2, no empty rows, no merged cells
- header names without spaces/special characters (Word turns them into field names)
- formatting for future input rows is set on the COLUMN, never on empty cells,
  otherwise Excel's used range grows and Word shows empty recipients
- formula cells get cached values injected (see inject_cached_values), because Word
  reads the stored values and never recalculates the workbook itself
"""
import re
import shutil
import sys
import zipfile
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

HEADERS = [
    # name, width, kind, comment
    ("Firma", 34, "input", "Name der Firma, genau wie in der Stellenanzeige."),
    ("Anrede", 10, "input", "Frau oder Herr auswählen (Pfeil rechts in der Zelle).\n"
                            "Leer lassen, wenn Sie keinen Namen kennen."),
    ("Vorname", 12, "input", "Vorname der Ansprechperson (kann leer bleiben)."),
    ("Nachname", 14, "input", "Nachname der Ansprechperson (kann leer bleiben)."),
    ("Strasse", 24, "input", "Straße und Hausnummer der Firma."),
    ("PLZ", 8, "input", "Postleitzahl. Die Spalte ist als TEXT formatiert,\n"
                        "damit eine 0 am Anfang nicht verschwindet."),
    ("Ort", 14, "input", "Ort der Firma."),
    ("Stelle", 22, "input", "Name der Stelle, z. B. Lagerhelfer.\n"
                            "Erscheint im Betreff: Bewerbung als ..."),
    ("Ansprechpartner", 24, "formula", "FORMEL – nicht überschreiben!\n"
                                       "Wird automatisch aus Anrede, Vorname und Nachname gebildet."),
    ("Briefanrede", 34, "formula", "FORMEL – nicht überschreiben!\n"
                                   "Frau → Sehr geehrte Frau ...\nHerr → Sehr geehrter Herr ...\n"
                                   "leer → Sehr geehrte Damen und Herren"),
    ("Quelle", 26, "tracker", "Wo haben Sie die Stelle gefunden?\n(Nur für Ihre Übersicht, nicht im Brief.)"),
    ("Beworben_am", 13, "tracker", "Datum der Bewerbung.\nHilft beim Nachweis für das Jobcenter."),
    ("Status", 12, "tracker", "offen / Einladung / Absage / Zusage"),
]

# Fictional sample data (no real companies, no real people).
ROWS = [
    ["Muster Logistik GmbH", "Frau", "Anna", "Petersen", "Am Musterhafen 12", "25813", "Husum",
     "Lagerhelfer", "Jobbörse Arbeitsagentur", date(2026, 11, 2), "offen"],
    ["Bäckerei Beispiel KG", "Herr", "Jens", "Hansen", "Beispielweg 3", "25821", "Bredstedt",
     "Verkäufer", "Aushang im Laden", None, "offen"],
    ["Pflegedienst Muster gGmbH", "Frau", "Maria", "Jensen", "Musterstraße 45", "25899", "Niebüll",
     "Pflegehelfer", "Firmen-Website", None, "offen"],
    ["Hotel Beispiel am Hafen", None, None, None, "Probeallee 7", "25832", "Tönning",
     "Küchenhilfe", "Zeitung", None, "offen"],
    ["Muster Gebäudeservice GmbH", "Herr", "Thomas", "Möller", "Beispielring 21", "24937", "Flensburg",
     "Hausmeister", "Jobbörse Arbeitsagentur", None, "offen"],
]

COL = {name: i + 1 for i, (name, *_rest) in enumerate(HEADERS)}
LETTER = {name: chr(ord("A") + i) for i, (name, *_rest) in enumerate(HEADERS)}

FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = {"input": "1F4E79", "formula": "7F7F7F", "tracker": "548235"}
formula_fill = PatternFill("solid", fgColor="EDEDED")


def ansprechpartner(anrede, vorname, nachname):
    return " ".join(p for p in (anrede or "", vorname or "", nachname or "") if p).strip()


def briefanrede(firma, anrede, nachname):
    if not firma:
        return ""
    if anrede == "Frau":
        return "Sehr geehrte Frau " + (nachname or "")
    if anrede == "Herr":
        return "Sehr geehrter Herr " + (nachname or "")
    return "Sehr geehrte Damen und Herren"


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Bewerbungen"

    for name, width, kind, note in HEADERS:
        c = ws.cell(row=1, column=COL[name], value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=head_fill[kind])
        c.alignment = Alignment(vertical="center")
        c.border = border
        c.comment = Comment(note, "EDV-Tag")
        c.comment.width = 260
        c.comment.height = 90
        dim = ws.column_dimensions[LETTER[name]]
        dim.width = width
        # Column-level formats apply to rows participants add later,
        # without creating empty cells that Word would read as recipients.
        dim.font = Font(name=FONT)
        if name == "PLZ":
            dim.number_format = "@"
        if name == "Beworben_am":
            dim.number_format = "DD.MM.YYYY"
    ws.row_dimensions[1].height = 20

    expected = {}
    for r, row in enumerate(ROWS, start=2):
        firma, anrede, vorname, nachname, strasse, plz, ort, stelle, quelle, beworben, status = row
        values = {
            "Firma": firma, "Anrede": anrede, "Vorname": vorname, "Nachname": nachname,
            "Strasse": strasse, "PLZ": plz, "Ort": ort, "Stelle": stelle,
            "Quelle": quelle, "Beworben_am": beworben, "Status": status,
        }
        for name, v in values.items():
            if v is not None:
                ws.cell(row=r, column=COL[name], value=v)
        a, b, c_, d = (f"{LETTER[n]}{r}" for n in ("Firma", "Anrede", "Vorname", "Nachname"))
        # German Excel shows these as GLÄTTEN / WENN with semicolons.
        ws.cell(row=r, column=COL["Ansprechpartner"], value=f'=TRIM({b}&" "&{c_}&" "&{d})')
        ws.cell(row=r, column=COL["Briefanrede"], value=(
            f'=IF({a}="","",IF({b}="Frau","Sehr geehrte Frau "&{d},'
            f'IF({b}="Herr","Sehr geehrter Herr "&{d},"Sehr geehrte Damen und Herren")))'))
        expected[f"{LETTER['Ansprechpartner']}{r}"] = ansprechpartner(anrede, vorname, nachname)
        expected[f"{LETTER['Briefanrede']}{r}"] = briefanrede(firma, anrede, nachname)

        for name in COL:
            cell = ws.cell(row=r, column=COL[name])
            cell.font = Font(name=FONT)
            cell.border = border
            if name in ("Ansprechpartner", "Briefanrede"):
                cell.fill = formula_fill
            if name == "PLZ":
                cell.number_format = "@"
            if name == "Beworben_am":
                cell.number_format = "DD.MM.YYYY"

    # Dropdowns. Data validation lives in its own XML element and does not create cells.
    dv_anrede = DataValidation(type="list", formula1='"Frau,Herr"', allow_blank=True,
                               showErrorMessage=True, errorTitle="Anrede",
                               error="Bitte Frau oder Herr auswählen – oder leer lassen.")
    dv_anrede.add(f"{LETTER['Anrede']}2:{LETTER['Anrede']}500")
    dv_status = DataValidation(type="list", formula1='"offen,Einladung,Absage,Zusage"',
                               allow_blank=True)
    dv_status.add(f"{LETTER['Status']}2:{LETTER['Status']}500")
    ws.add_data_validation(dv_anrede)
    ws.add_data_validation(dv_status)

    ws.freeze_panes = "B2"
    # Printable as proof of applications: landscape, one page wide (no print titles:
    # they create a defined name that shows up as an extra table in Word's picker).
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.zoomScale = 110
    wb.save(OUT)
    return expected


def inject_cached_values(path, expected):
    """openpyxl writes formulas without results; Word's Excel reader only sees stored results."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                for ref, value in expected.items():
                    pattern = re.compile(rf'<c r="{ref}"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>')
                    m = pattern.search(xml)
                    if not m:
                        raise SystemExit(f"formula cell {ref} not found in sheet XML")
                    attrs = re.sub(r'\s*t="[^"]*"', "", m.group(1))
                    esc = value.replace("&", "&amp;").replace("<", "&lt;")
                    xml = xml[:m.start()] + f'<c r="{ref}"{attrs} t="str"><f>{m.group(2)}</f><v>{esc}</v></c>' + xml[m.end():]
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


if __name__ == "__main__":
    exp = build()
    inject_cached_values(OUT, exp)
    for k, v in exp.items():
        print(k, "=>", repr(v))
